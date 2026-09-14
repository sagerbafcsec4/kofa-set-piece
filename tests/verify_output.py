#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""出来上がったセットプレーExcelを機械で照合する台本。

使い方:
    python tests\\verify_output.py                     tests/output の最新xlsxを照合
    python tests\\verify_output.py <出力xlsx>          指定ファイルを照合
    python tests\\verify_output.py <出力xlsx> --dict docs/opta-names-draft.csv
                                                     辞書をGoogleシートでなくCSVから取る

判定（すべて ✔ で PASS）:
  (a) 値: Opta の2ファイルと全セル一致（Team は辞書適用後）・表題・見出し・シート名・結合
  (b) 並び: セットプレー降順 → Total 降順 → 名前昇順
  (c) 体裁: フォント・太字・色・塗り・配置・罫線・列幅・行高・数値書式
  (d) 数式なし・Excelエラー値なし・Total合計が元の Goals From Set Piece 合計と一致
"""
import csv
import glob
import io
import os
import re
import sys
import unicodedata
import urllib.request

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import paths  # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

OPTA_COLS = {
    "team": "team name", "corner": "goals from corner", "direct": "goals from direct freekicks",
    "indirect": "goals from indirect freekicks", "penalty": "goals from penalties",
    "throwIn": "goals from throw in", "setPiece": "goals from set piece", "total": "total",
    "setPiecePct": "goals from set piece %",
}
OUT_HEADERS = ["Team", "Total", "Penalties", "Corners", "Direct\nFreekicks", "Indirect\nFreekicks", "Throw In", "Goals From\nSet Piece ％"]
COL_WIDTHS = [25, 13, 13, 13, 13, 13, 13, 13]
ROW_H = {"title": 19.5, "head": 30, "body": 17.25, "gap": 18, "title2": 18.6}
FONT = "MS UI Gothic"
TITLE_FILL, HEAD_FILL, WHITE, BLACK = "FFE7E6E6", "FF000000", "FFFFFFFF", "FF000000"
ERR_RE = re.compile(r"#(REF!|DIV/0!|VALUE!|NAME\?|N/A|NUM!|NULL!)")


def norm_header(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def norm_key(s):
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(s or ""))).strip().lower()


def read_opta(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Sheet"]
    header = {norm_header(c.value): i for i, c in enumerate(ws[1]) if c.value is not None}
    idx = {k: header[v] for k, v in OPTA_COLS.items()}
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        team = str(r[idx["team"]] or "").strip()
        if not team:
            continue
        rec = {"team": team}
        for k in ("setPiece", "penalty", "corner", "direct", "indirect", "throwIn", "total"):
            rec[k] = float(r[idx[k]])
        rec["setPiecePct"] = round(float(r[idx["setPiecePct"]]) * 100) / 10000
        rows.append(rec)
    src_sum = sum(float(r[idx["setPiece"]]) for r in ws.iter_rows(min_row=2, values_only=True) if r[idx["team"]])
    return rows, src_sum


BUILTIN_CSV = os.path.join(paths.REPO, "docs", "opta-names-draft.csv")


def load_dict(league, dict_csv=None):
    """内蔵表（docs/opta-names-draft.csv＝index.html の BUILTIN_DICT と同内容）を土台に、
    共有シートに「Opta表記」列があればそれで上書きする（アプリと同じ優先順）。"""
    base, _ = _load_dict_from_text(open(BUILTIN_CSV, encoding="utf-8-sig").read(), league)
    if dict_csv:
        text = open(dict_csv, encoding="utf-8-sig").read()
    else:
        try:
            text = urllib.request.urlopen(paths.SHEET_CSV, timeout=30).read().decode("utf-8-sig")
        except Exception:
            return base, False
    sheet, has_col = _load_dict_from_text(text, league)
    base.update(sheet)
    return base, has_col


def _load_dict_from_text(text, league):
    rows = [r for r in csv.reader(io.StringIO(text)) if any(c.strip() for c in r)]
    header = [h.strip() for h in rows[0]]

    def find(*keys):
        for i, h in enumerate(header):
            if any(k in h for k in keys):
                return i
        return -1

    i_name, i_league, i_opta = find("チーム名"), find("所属リーグ", "リーグ"), find("Opta表記")
    if i_opta < 0:
        i_opta = next((i for i, h in enumerate(header) if re.search("opta", h, re.I)), -1)
    d = {}
    for r in rows[1:]:
        if len(r) <= max(i_name, i_league):
            continue
        if r[i_league].strip() != league or not r[i_name].strip():
            continue
        if i_opta < 0 or i_opta >= len(r):
            continue
        for part in re.split(r"[;；]", r[i_opta]):
            if part.strip():
                d[norm_key(part)] = r[i_name].strip()
    return d, i_opta >= 0


def expected_rows(opta_rows, dic):
    out = []
    for r in opta_rows:
        rec = dict(r)
        rec["name"] = dic.get(norm_key(r["team"]), r["team"])
        out.append(rec)
    out.sort(key=lambda x: (-x["setPiece"], -x["total"], x["name"]))
    return out


def border_style(cell, side):
    s = getattr(cell.border, side)
    return s.style if s is not None else None


def fill_rgb(cell):
    return cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type == "solid" else None


def font_rgb(cell):
    try:
        return cell.font.color.rgb if cell.font.color is not None else None
    except Exception:
        return "theme"


class Report:
    def __init__(self):
        self.items = []

    def check(self, ok, label, detail=""):
        self.items.append((bool(ok), label, detail))
        return ok

    def summary(self):
        n_ok = sum(1 for ok, _, _ in self.items if ok)
        lines = ["%s %s%s" % ("✔" if ok else "✘", label, ("  — " + detail) if detail else "") for ok, label, detail in self.items]
        lines.append("PASS %d/%d" % (n_ok, len(self.items)))
        return "\n".join(lines), n_ok == len(self.items)


def check_table(ws, rep, start_row, title, rows, is_second, tag):
    r0, r1 = start_row, start_row + 1
    n = len(rows)
    last = r1 + n
    # 表題
    rep.check(ws.cell(r0, 1).value == title, f"{tag} 表題の文言", f"{ws.cell(r0, 1).value!r}")
    rep.check(f"A{r0}:H{r0}" in {str(m) for m in ws.merged_cells.ranges}, f"{tag} 表題が A{r0}:H{r0} で結合")
    rep.check(abs((ws.row_dimensions[r0].height or 0) - (ROW_H["title2"] if is_second else ROW_H["title"])) < 0.01, f"{tag} 表題の行高", str(ws.row_dimensions[r0].height))
    # openpyxl は結合範囲の2番目以降を MergedCell（文字・塗り・配置なし）として読み、
    # 外周の罫線だけを各セルに合成する（先頭セルには四方 medium が集まる）。前回の完成品も同じ見え方だった。
    bad = []
    for c in range(1, 9):
        cell = ws.cell(r0, c)
        if c == 1:
            exp = dict(font=FONT, size=12, bold=True, color=BLACK, fill=TITLE_FILL, h="center", v="center", shrink=True,
                       top="medium", bottom="medium", left="medium", right="medium")
            got = dict(font=cell.font.name, size=cell.font.size, bold=bool(cell.font.bold), color=font_rgb(cell), fill=fill_rgb(cell),
                       h=cell.alignment.horizontal, v=cell.alignment.vertical, shrink=bool(cell.alignment.shrink_to_fit),
                       top=border_style(cell, "top"), bottom=border_style(cell, "bottom"), left=border_style(cell, "left"), right=border_style(cell, "right"))
        else:
            exp = dict(top="medium", bottom="medium", right="medium" if c == 8 else None)
            got = dict(top=border_style(cell, "top"), bottom=border_style(cell, "bottom"), right=border_style(cell, "right"))
        diff = {k: (exp[k], got[k]) for k in exp if exp[k] != got[k]}
        if diff:
            bad.append(f"{cell.coordinate}:{diff}")
    rep.check(not bad, f"{tag} 表題行の体裁（先頭セル全項目＋外周罫線）", "; ".join(bad)[:300])
    # 見出し
    rep.check([ws.cell(r1, c).value for c in range(1, 9)] == OUT_HEADERS, f"{tag} 見出し8個", str([ws.cell(r1, c).value for c in range(1, 9)]))
    rep.check(abs((ws.row_dimensions[r1].height or 0) - ROW_H["head"]) < 0.01, f"{tag} 見出しの行高", str(ws.row_dimensions[r1].height))
    bad = []
    for c in range(1, 9):
        cell = ws.cell(r1, c)
        exp = dict(font=FONT, size=12 if c <= 4 else 11, bold=True, color=WHITE, fill=HEAD_FILL, h="center", v="center", wrap=True,
                   top="medium", bottom="thin", left="medium" if c == 1 else "thin", right="medium" if c == 8 else "thin",
                   fmt="@" if c == 8 else "General")
        got = dict(font=cell.font.name, size=cell.font.size, bold=bool(cell.font.bold), color=font_rgb(cell), fill=fill_rgb(cell),
                   h=cell.alignment.horizontal, v=cell.alignment.vertical, wrap=bool(cell.alignment.wrap_text),
                   top=border_style(cell, "top"), bottom=border_style(cell, "bottom"), left=border_style(cell, "left"), right=border_style(cell, "right"),
                   fmt=cell.number_format)
        diff = {k: (exp[k], got[k]) for k in exp if exp[k] != got[k]}
        if diff:
            bad.append(f"{cell.coordinate}:{diff}")
    rep.check(not bad, f"{tag} 見出し行の体裁（8セル）", "; ".join(bad)[:300])
    # 本文の値
    bad = []
    for i, rec in enumerate(rows):
        r = r1 + 1 + i
        got = [ws.cell(r, c).value for c in range(1, 9)]
        exp = [rec["name"], rec["setPiece"], rec["penalty"], rec["corner"], rec["direct"], rec["indirect"], rec["throwIn"], rec["setPiecePct"]]
        if str(got[0]) != exp[0]:
            bad.append(f"A{r}: {got[0]!r} != {exp[0]!r}")
        for c in range(1, 8):
            try:
                if abs(float(got[c]) - float(exp[c])) > 1e-9:
                    bad.append(f"{ws.cell(r, c + 1).coordinate}: {got[c]!r} != {exp[c]!r}")
            except (TypeError, ValueError):
                bad.append(f"{ws.cell(r, c + 1).coordinate}: 数値でない {got[c]!r}")
    rep.check(not bad, f"{tag} 本文の値と並び（{n}行×8列）", "; ".join(bad)[:400])
    # 本文の体裁
    bad = []
    for i in range(n):
        r = r1 + 1 + i
        is_last = r == last
        if abs((ws.row_dimensions[r].height or 0) - ROW_H["body"]) > 0.01:
            bad.append(f"行{r}の高さ {ws.row_dimensions[r].height}")
        for c in range(1, 9):
            cell = ws.cell(r, c)
            exp = dict(font=FONT, size=12 if c <= 4 else 11, bold=False, fill=None, h="center", v="center", shrink=True,
                       top="thin", bottom="medium" if is_last else "thin", left="medium" if c == 1 else "thin", right="medium" if c == 8 else "thin",
                       fmt="0.00%" if c == 8 else "General")
            got = dict(font=cell.font.name, size=cell.font.size, bold=bool(cell.font.bold), fill=fill_rgb(cell),
                       h=cell.alignment.horizontal, v=cell.alignment.vertical, shrink=bool(cell.alignment.shrink_to_fit),
                       top=border_style(cell, "top"), bottom=border_style(cell, "bottom"), left=border_style(cell, "left"), right=border_style(cell, "right"),
                       fmt=cell.number_format)
            diff = {k: (exp[k], got[k]) for k in exp if exp[k] != got[k]}
            if diff:
                bad.append(f"{cell.coordinate}:{diff}")
    rep.check(not bad, f"{tag} 本文の体裁（{n * 8}セル）", "; ".join(bad)[:400])
    return last


def verify(out_path, dict_csv=None, league=paths.LEAGUE, matchday=paths.MATCHDAY, goals_path=paths.GOALS, conceded_path=paths.CONCEDED):
    rep = Report()
    wb = openpyxl.load_workbook(out_path)
    rep.check(wb.sheetnames == [f"{matchday}節用セットプレー"], "シート名", str(wb.sheetnames))
    ws = wb.worksheets[0]
    goals, goals_sum = read_opta(goals_path)
    conceded, conceded_sum = read_opta(conceded_path)
    dic, has_col = load_dict(league, dict_csv)
    rep.check(True, "辞書", f"{league} 対応 {len(dic)}件（内蔵表" + ("＋共有シートのOpta表記列）" if has_col else "のみ・シートに列なし）"))
    eg, ec = expected_rows(goals, dic), expected_rows(conceded, dic)
    # 列幅
    bad = [f"{chr(64 + i + 1)}={ws.column_dimensions[chr(64 + i + 1)].width}" for i, w in enumerate(COL_WIDTHS)
           if abs((ws.column_dimensions[chr(64 + i + 1)].width or 0) - w) > 0.1]
    rep.check(not bad, "列幅 A〜H", "; ".join(bad))
    end1 = check_table(ws, rep, 1, f"セットプレーからの得点数 (Opta) ※第{matchday}節終了時", eg, False, "得点表")
    gap = end1 + 1
    rep.check(abs((ws.row_dimensions[gap].height or 0) - ROW_H["gap"]) < 0.01 and all(ws.cell(gap, c).value is None for c in range(1, 9)), "空き行", f"行{gap} 高さ {ws.row_dimensions[gap].height}")
    end2 = check_table(ws, rep, gap + 1, f"セットプレーからの失点数 (Opta) ※第{matchday}節終了時", ec, True, "失点表")
    rep.check(ws.max_row == end2, "余計な行がない", f"max_row={ws.max_row} 期待={end2}")
    # 数式・エラー値
    bad = []
    for row in ws.iter_rows():
        for c in row:
            if c.data_type == "f":
                bad.append(f"{c.coordinate} 数式")
            elif isinstance(c.value, str) and ERR_RE.search(c.value):
                bad.append(f"{c.coordinate} {c.value}")
    rep.check(not bad, "数式なし・Excelエラー値なし", "; ".join(bad)[:200])
    # Total 合計
    sg = sum(float(ws.cell(r, 2).value or 0) for r in range(3, end1 + 1))
    sc = sum(float(ws.cell(r, 2).value or 0) for r in range(gap + 3, end2 + 1))
    rep.check(sg == goals_sum and sc == conceded_sum, "Total 合計が元の Goals From Set Piece 合計と一致", f"得点 {sg}（元 {goals_sum}）／失点 {sc}（元 {conceded_sum}）")
    text, ok = rep.summary()
    return ok, text


def main():
    args = sys.argv[1:]
    dict_csv = None
    if "--dict" in args:
        dict_csv = args[args.index("--dict") + 1]
        del args[args.index("--dict"):args.index("--dict") + 2]
    if args:
        out = args[0]
    else:
        cands = sorted(glob.glob(os.path.join(paths.OUT_DIR, "*.xlsx")), key=os.path.getmtime)
        if not cands:
            print("照合する出力xlsxが tests/output にありません")
            return 2
        out = cands[-1]
    print("照合対象:", out)
    ok, text = verify(out, dict_csv)
    print(text)
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
